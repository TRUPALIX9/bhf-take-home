"""Tests for po_extract.py, grouped by section: parse, validate, write, main.

Run from python-po-parser/: python3 -m pytest -q
They use the real sample PO and template from input/, and write only to tmp_path.
"""

import json
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

import po_extract as pe

INPUT = Path(__file__).resolve().parent.parent / "input"
SAMPLE = INPUT / "purchase_order_sample.txt"
TEMPLATE = INPUT / "output_template_blank.xlsx"
# The template's own header row, read once. Tests compare against this
# instead of a hard-coded list, so they follow the file the brief supplied.
TEMPLATE_HEADER = [cell.value for cell in load_workbook(TEMPLATE)["Sheet1"][1]]


@pytest.fixture(scope="module")
def po():
    """The sample PO, parsed once and shared by every test in this file."""
    return pe.parse(SAMPLE.read_text(encoding="utf-8"))


def write_and_open(po, tmp_path):
    """Write the workbook to tmp_path and return its only sheet."""
    out = tmp_path / "out.xlsx"
    pe.write_xlsx(po, SAMPLE.name, TEMPLATE, out)
    return load_workbook(out)["Sheet1"]


def row_as_dict(sheet, row):
    """One sheet row as {template column name: cell value}."""
    return dict(zip(TEMPLATE_HEADER, [cell.value for cell in sheet[row]]))


# ---------------------------------------------------------------- parse


def test_header_fields(po):
    # Every labeled header value, plus both dates turned into real dates.
    header = po.header
    assert header.buyer == "610 JORDAN MILLER"
    assert header.ship_terms == "FOB Ningbo,CHN"
    assert header.po_number == "TX9K2QP"
    assert header.ref_master_po == "TX9H1LM"
    assert header.po_date == date(2026, 3, 14)
    assert header.ship_date == date(2026, 5, 2)


def test_address_blocks(po):
    # VENDOR is a single block; SHIP TO and BILL TO share lines and are split by column.
    header = po.header
    assert header.vendor == (
        "50201 HARBOR POINT TRADING LTD",
        "C/O MERIDIAN SOURCING CO INC",
        "410 COMMERCE AVE",
        "EDISON NJ 08837",
        "PHONE : 9085551040",
        "FAX : 9085551099",
    )
    assert header.ship_to == (
        "331 DISTRIBUTION/NORTHEAST",
        "C/O SWIFTLINE EXPRESS",
        "1450 CROSSROADS BLVD",
        "EDISON, NJ 08817",
        "PHONE : 9085551088",
    )
    assert header.bill_to == (
        "331 MERIDIAN SOURCING CO. INC.",
        "C/O NORTHSTAR RETAIL GROUP - DEPT OPS",
        "PO BOX 4021",
        "EDISON, NJ 08837",
    )


def test_twelve_items_each_with_its_upc(po):
    # The brief's evaluation names the UPC line explicitly: each item must get
    # the UPC printed under it, with leading zeros kept.
    assert len(po.items) == 12
    assert [item.sku for item in po.items][:2] == ["71097188", "71097195"]
    assert po.items[0].upc == "0091120032129"
    assert po.items[-1].upc == "0091120032426"
    assert all(len(item.upc) == 13 for item in po.items)


def test_first_row_and_edge_cases(po):
    # Every column of row 1, then two edge cases: a description containing
    # "F/Q" (row 8) and the widest EXT COST, 124801.600 (row 5).
    first = po.items[0]
    assert (first.dept, first.vendor_part) == ("210", "04212")
    assert first.description == "NORDPEAK ASHER T CMFST GRY"
    assert (first.retail, first.cost, first.ext_cost) == (
        Decimal("149.99"),
        Decimal("49.400"),
        Decimal("6718.400"),
    )
    assert (first.cartons, first.case_pack, first.ext_qty) == (68, 2, 136)
    assert (first.cube, first.kilograms) == (Decimal("0.0991"), Decimal("7.72"))
    assert po.items[7].description == "NORDPEAK ASHER STRIPE F/Q DVT GRAY"
    assert po.items[4].ext_cost == Decimal("124801.600")


def test_missing_upc_raises_naming_the_sku():
    # A missing UPC must stop the run, not shift every later UPC up one row.
    text = SAMPLE.read_text(encoding="utf-8").replace("    0091120032129\n", "")
    with pytest.raises(ValueError, match="71097188"):
        pe.parse(text)


def test_unreadable_item_row_raises_instead_of_being_dropped():
    # A row with a missing number no longer matches the item regex.
    # The error must name the line instead of silently writing 11 rows.
    text = SAMPLE.read_text(encoding="utf-8").replace("  68     2      136", "  68     2")
    with pytest.raises(ValueError, match="line 25 is not an item row"):
        pe.parse(text)


# ---------------------------------------------------------------- validate


def test_sample_has_no_flags(po):
    # The sample is internally consistent, so a clean run has zero flags.
    assert pe.all_flags(po) == []
    assert po.totals == pe.Totals(Decimal("447455.120"), 3475, 6950, Decimal("119.18"))


def test_tampered_rows_are_flagged(po):
    # A wrong EXT QTY also breaks EXT Cost (COST x EXT QTY); a wrong EXT Cost alone flags one.
    bad_qty = replace(po.items[0], ext_qty=135)
    assert [flag.column for flag in pe.check_item(bad_qty)] == ["EXT QTY", "EXT Cost"]
    bad_cost = replace(po.items[0], ext_cost=Decimal("6718.401"))
    assert [flag.column for flag in pe.check_item(bad_cost)] == ["EXT Cost"]


def test_tampered_text_is_flagged_in_rows_and_totals():
    # Same check end to end from the text: one changed number shows up on the
    # row and in the TOTALS comparison.
    text = SAMPLE.read_text(encoding="utf-8").replace(
        "   673     2     1346", "   673     2     1347"
    )
    po = pe.parse(text)
    assert [flag.column for flag in pe.check_item(po.items[1])] == ["EXT QTY", "EXT Cost"]
    assert any(message.startswith("TOTALS EXT QTY") for message in pe.check_totals(po))


# ---------------------------------------------------------------- write


def test_header_row_and_sheets_match_template(po, tmp_path):
    # The output keeps the template's single sheet and header row, with 12 item rows.
    out = tmp_path / "out.xlsx"
    pe.write_xlsx(po, SAMPLE.name, TEMPLATE, out)
    workbook = load_workbook(out)
    assert workbook.sheetnames == ["Sheet1"]
    sheet = workbook["Sheet1"]
    assert [cell.value for cell in sheet[1]] == TEMPLATE_HEADER
    assert sheet.max_column == len(TEMPLATE_HEADER)
    assert sheet.max_row == 13


def test_row_values(po, tmp_path):
    # Values read back from the saved file. Header fields repeat on every row.
    sheet = write_and_open(po, tmp_path)
    first = row_as_dict(sheet, 2)
    assert first["FILENAME"] == "purchase_order_sample.txt"
    assert first["PO#"] == "TX9K2QP"
    assert first["SHIP DATE"] == datetime(2026, 5, 2)
    assert first["VENDOR"].splitlines()[0] == "50201 HARBOR POINT TRADING LTD"
    assert first["Cost"] == 49.4 and first["EXT Cost"] == 6718.4
    assert first["EXT QTY"] == 136
    assert first["Total Cost"] == 447455.12
    last = row_as_dict(sheet, 13)
    assert last["BUYER"] == "610 JORDAN MILLER" and last["SKU"] == "71293283"


def test_leading_zeros_preserved(po, tmp_path):
    # Code columns are strings in Text ("@") format, so Excel keeps leading zeros.
    sheet = write_and_open(po, tmp_path)
    first = row_as_dict(sheet, 2)
    assert first["UPC"] == "0091120032129"
    assert first["Vendor Part #"] == "04212"
    assert first["SKU"] == "71097188" and first["Dept"] == "210"
    for column in ["Dept", "SKU", "UPC", "Vendor Part #", "PO#"]:
        cell = sheet.cell(2, TEMPLATE_HEADER.index(column) + 1)
        assert cell.number_format == "@"


def test_flag_lands_on_the_failing_cell(po, tmp_path):
    # A bad row gets the red fill and comment on the exact cell (T = EXT QTY,
    # Q = EXT Cost), and the next row stays clean.
    bad_items = (replace(po.items[0], ext_qty=135),) + po.items[1:]
    sheet = write_and_open(replace(po, items=bad_items), tmp_path)
    ext_qty = sheet["T2"]
    assert sheet["T1"].value == "EXT QTY"
    assert ext_qty.fill.fgColor.rgb.endswith("FFC7CE")
    assert "135" in ext_qty.comment.text
    assert sheet["Q2"].comment is not None  # EXT Cost is off too
    assert sheet["T3"].comment is None


def test_unknown_template_column_raises(po, tmp_path):
    # A template column with no mapping fails loudly instead of being left blank.
    template = tmp_path / "template.xlsx"
    workbook = load_workbook(TEMPLATE)
    sheet = workbook["Sheet1"]
    sheet.cell(1, sheet.max_column + 1, "Color")
    workbook.save(template)
    with pytest.raises(ValueError, match="Color"):
        pe.write_xlsx(po, SAMPLE.name, template, tmp_path / "out.xlsx")


# ---------------------------------------------------------------- main


def run_main(tmp_path, input_path=SAMPLE, *extra):
    """Call main() as the command line would, writing into tmp_path."""
    out = tmp_path / "out.xlsx"
    args = ["--input", str(input_path), "--template", str(TEMPLATE), "--output", str(out)]
    return pe.main(args + list(extra))


def test_main_exit_code_zero_on_sample(tmp_path, capsys):
    # Clean run: exit 0, spreadsheet written, and no JSON unless --json is given.
    assert run_main(tmp_path) == 0
    assert (tmp_path / "out.xlsx").exists()
    assert not (tmp_path / "TX9K2QP.json").exists()
    assert "Flags: 0" in capsys.readouterr().out


def test_main_exit_code_one_when_flagged(tmp_path, capsys):
    # Exit 1 when anything is flagged, so a script or CI can catch it.
    tampered = tmp_path / "tampered.txt"
    tampered.write_text(
        SAMPLE.read_text(encoding="utf-8").replace("6718.400", "6718.401"),
        encoding="utf-8",
    )
    assert run_main(tmp_path, tampered) == 1
    assert "FLAG SKU 71097188 EXT Cost" in capsys.readouterr().out


def test_json_export(tmp_path):
    # --json writes the Next.js seed. The keys must match the app's Order and LineItem types.
    assert run_main(tmp_path, SAMPLE, "--json") == 0
    order = json.loads((tmp_path / "TX9K2QP.json").read_text(encoding="utf-8"))
    assert set(order) == {
        "poNumber", "refMasterPo", "poDate", "shipDate", "buyer", "shipTerms",
        "vendor", "shipTo", "billTo", "totalCost", "status", "lines",
    }
    assert order["poNumber"] == "TX9K2QP"
    assert order["poDate"] == "2026-03-14" and order["shipDate"] == "2026-05-02"
    assert order["status"] == "pending"
    assert order["totalCost"] == 447455.12
    assert len(order["lines"]) == 12
    first = order["lines"][0]
    assert set(first) == {
        "dept", "sku", "upc", "vendorPart", "description", "retail", "cost",
        "extCost", "cartons", "casePack", "extQty", "cube", "kilograms",
    }
    assert first["upc"] == "0091120032129" and first["vendorPart"] == "04212"
    assert first["extCost"] == 6718.4 and first["extQty"] == 136
